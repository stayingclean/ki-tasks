from pathlib import Path

from docx import Document

from platzhalter.app import Api, lade_html


def speichere(pfad: Path, text: str) -> None:
    doc = Document()
    doc.add_paragraph(text)
    doc.save(str(pfad))


def test_api_suche_und_ersetze(tmp_path: Path):
    speichere(tmp_path / "a.docx", "[AA] [BB]")
    speichere(tmp_path / "b.docx", "[BB]")
    api = Api(tmp_path)

    assert api.hole_ordner() == str(tmp_path)
    ergebnis = api.suche()
    assert ergebnis["dateien"] == ["a.docx", "b.docx"]
    assert [p["name"] for p in ergebnis["platzhalter"]] == ["BB", "AA"]

    resultate = api.ersetze({"AA": "eins", "BB": ""})
    assert resultate == [
        {"datei": "a.docx", "ersetzt": 1, "fehler": None},
        {"datei": "b.docx", "ersetzt": 0, "fehler": None},
    ]


def test_lade_html_bettet_logo_ein():
    html = lade_html()
    assert "{{LOGO}}" not in html
    assert 'src="data:image/png;base64,iVBORw0KGgo' in html
    assert len(html) > 20000  # das eingebettete Logo macht den Grossteil aus


def test_api_ersetzt_nichts_im_verlauf_ordner(tmp_path: Path):
    (tmp_path / "verlauf").mkdir()
    speichere(tmp_path / "verlauf" / "2026-09-18_a.docx", "[AA]")
    api = Api(tmp_path / "verlauf")

    ergebnis = api.suche()
    assert ergebnis["warnung"]
    assert ergebnis["dateien"] == []

    assert api.ersetze({"AA": "eins"}) == []
    alt = Document(str(tmp_path / "verlauf" / "2026-09-18_a.docx"))
    assert alt.paragraphs[0].text == "[AA]"


def test_api_ersetze_meldet_fehler_pro_datei(tmp_path: Path):
    speichere(tmp_path / "a.docx", "[AA]")
    api = Api(tmp_path)
    api.suche()
    (tmp_path / "a.docx").write_bytes(b"jetzt kaputt")
    resultate = api.ersetze({"AA": "x"})
    assert resultate[0]["datei"] == "a.docx"
    assert resultate[0]["ersetzt"] == 0
    assert resultate[0]["fehler"]
