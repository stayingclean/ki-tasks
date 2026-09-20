from pathlib import Path

from docx import Document

from platzhalter.docx_ersetzen import ersetze_im_absatz, ersetze_in_datei
from platzhalter.dokument import absatz_text, alle_absaetze


def absatz_mit_runs(*teile: str):
    doc = Document()
    p = doc.add_paragraph()
    for t in teile:
        p.add_run(t)
    return p


def test_platzhalter_in_einem_run():
    p = absatz_mit_runs("Wohnhaft an [ADRESSE].")
    assert ersetze_im_absatz(p._p, {"ADRESSE": "Musterweg 1"}) == 1
    assert p.text == "Wohnhaft an Musterweg 1."


def test_platzhalter_ueber_drei_runs_verteilt():
    p = absatz_mit_runs("Hallo [AD", "RES", "SE], danke")
    assert ersetze_im_absatz(p._p, {"ADRESSE": "Musterweg 1"}) == 1
    assert p.text == "Hallo Musterweg 1, danke"
    assert [r.text for r in p.runs] == ["Hallo Musterweg 1", "", ", danke"]


def test_zwei_vorkommen_im_selben_absatz():
    p = absatz_mit_runs("[AA] und [AA] und [BB]")
    assert ersetze_im_absatz(p._p, {"AA": "eins", "BB": "zwei"}) == 3
    assert p.text == "eins und eins und zwei"


def test_leerer_oder_fehlender_wert_bleibt_stehen():
    p = absatz_mit_runs("[AA] [BB] [CC]")
    assert ersetze_im_absatz(p._p, {"AA": "", "BB": "x"}) == 1
    assert p.text == "[AA] x [CC]"


def test_formatierung_des_ersten_runs_bleibt():
    p = absatz_mit_runs("Name: ", "[NA", "ME]")
    p.runs[1].bold = True
    ersetze_im_absatz(p._p, {"NAME": "Muster"})
    assert p.runs[0].text == "Name: " and p.runs[0].bold is None
    assert p.runs[1].text == "Muster" and p.runs[1].bold is True


def test_tab_im_run_bleibt_erhalten():
    p = absatz_mit_runs("Name:\t[NAME]")
    ersetze_im_absatz(p._p, {"NAME": "Muster"})
    assert p.text == "Name:\tMuster"


def test_ersetze_in_datei_tabelle_kopfzeile_und_sicherung(tmp_path: Path):
    pfad = tmp_path / "brief.docx"
    doc = Document()
    doc.add_paragraph("Sehr geehrte [ANREDE]")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "IBAN: [IBAN]"
    doc.sections[0].header.paragraphs[0].text = "[NAME], [ADRESSE]"
    doc.save(str(pfad))

    n = ersetze_in_datei(pfad, {"ANREDE": "Frau Muster", "IBAN": "CH00", "NAME": "M", "ADRESSE": ""})

    assert n == 3
    texte = [absatz_text(p) for p in alle_absaetze(Document(str(pfad)))]
    assert "Sehr geehrte Frau Muster" in texte
    assert "IBAN: CH00" in texte
    assert "M, [ADRESSE]" in texte
    assert (tmp_path / "brief.docx.bak").exists()
    assert not (tmp_path / "brief.docx.tmp").exists()
    sicherung = [absatz_text(p) for p in alle_absaetze(Document(str(tmp_path / "brief.docx.bak")))]
    assert "Sehr geehrte [ANREDE]" in sicherung


def test_ersetze_in_datei_ohne_treffer_laesst_datei_unveraendert(tmp_path: Path):
    pfad = tmp_path / "x.docx"
    doc = Document()
    doc.add_paragraph("[AA]")
    doc.save(str(pfad))
    vorher = pfad.read_bytes()

    assert ersetze_in_datei(pfad, {"BB": "x", "AA": ""}) == 0
    assert pfad.read_bytes() == vorher
    assert not (tmp_path / "x.docx.bak").exists()
