from pathlib import Path

from docx import Document

from platzhalter.dokument import MUSTER, absatz_text, alle_absaetze
from platzhalter.suche import durchsuche, finde_dateien

from formular import formular, ohne_formular


def speichere(pfad: Path, *absaetze: str) -> Path:
    doc = Document()
    for a in absaetze:
        doc.add_paragraph(a)
    pfad.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(pfad))
    return pfad


def test_muster_findet_platzhalter():
    assert MUSTER.findall("Hallo [ADRESSE], [Ort] und [AHV]") == ["ADRESSE", "Ort", "AHV"]


def test_muster_erlaubt_umlaute():
    assert MUSTER.findall("[Grösse] [MASSNAHMEN] [Änderung]") == ["Grösse", "MASSNAHMEN", "Änderung"]


def test_muster_nimmt_ziffern_und_zeichen_neben_dem_wort():
    text = "[AHV-Nummer] [PLZ/Ort] [Betrag 2026] [Datum: 1.1.]"
    assert MUSTER.findall(text) == ["AHV-Nummer", "PLZ/Ort", "Betrag 2026", "Datum: 1.1."]


def test_muster_laesst_quellenverweise_liegen():
    text = "Studie [1] und [1, 2] sowie [12] und [S. 12] und [3a] und []"
    assert MUSTER.findall(text) == []


def test_muster_ignoriert_verschachtelte_und_umbrueche():
    assert MUSTER.findall("[[xy]] [ab\ncd] [ab\tcd]") == ["xy"]


def test_muster_begrenzt_laenge():
    assert MUSTER.findall("[" + "x" * 80 + "]") == ["x" * 80]
    assert MUSTER.findall("[" + "x" * 81 + "]") == []


def test_finde_dateien_rekursiv_und_ueberspringt_sperr_und_sicherungsdateien(tmp_path):
    speichere(tmp_path / "a.docx", "x")
    speichere(tmp_path / "unter" / "tief" / "b.docx", "x")
    speichere(tmp_path / "~$a.docx", "x")
    ohne_formular(tmp_path / "c.pdf")
    ohne_formular(tmp_path / "e.PDF")
    (tmp_path / "a.docx.bak").write_bytes(b"egal")
    (tmp_path / "c.pdf.bak").write_bytes(b"egal")
    (tmp_path / "notiz.txt").write_text("egal")
    gefunden = finde_dateien(tmp_path)
    assert [p.relative_to(tmp_path).as_posix() for p in gefunden] == [
        "a.docx", "c.pdf", "e.PDF", "unter/tief/b.docx",
    ]


def test_alle_absaetze_umfasst_tabelle_und_kopfzeile(tmp_path):
    doc = Document()
    doc.add_paragraph("Body [AA]")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Zelle [BB]"
    doc.sections[0].header.paragraphs[0].text = "Kopf [CC]"
    doc.sections[0].footer.paragraphs[0].text = "Fuss [DD]"
    texte = [absatz_text(p) for p in alle_absaetze(doc)]
    assert "Body [AA]" in texte
    assert "Zelle [BB]" in texte
    assert "Kopf [CC]" in texte
    assert "Fuss [DD]" in texte


def test_durchsuche_zaehlt_und_sortiert(tmp_path):
    speichere(tmp_path / "eins.docx", "[ADRESSE] und [ADRESSE]", "[Ort]")
    speichere(tmp_path / "sub" / "zwei.docx", "[ADRESSE]", "[AHV]")
    e = durchsuche(tmp_path)
    assert e.ordner == str(tmp_path)
    assert e.dateien == ["eins.docx", "sub/zwei.docx"]
    assert e.fehler == []
    assert [(f.name, f.anzahl, f.dateien) for f in e.platzhalter] == [
        ("ADRESSE", 3, ["eins.docx", "sub/zwei.docx"]),
        ("AHV", 1, ["sub/zwei.docx"]),
        ("Ort", 1, ["eins.docx"]),
    ]


def test_durchsuche_meldet_kaputte_datei(tmp_path):
    (tmp_path / "kaputt.docx").write_bytes(b"kein zip")
    speichere(tmp_path / "ok.docx", "[XX]")
    e = durchsuche(tmp_path)
    assert e.dateien == ["ok.docx"]
    assert len(e.fehler) == 1 and e.fehler[0].startswith("kaputt.docx: ")
    assert [f.name for f in e.platzhalter] == ["XX"]


def test_finde_dateien_ueberspringt_verlauf(tmp_path):
    speichere(tmp_path / "motivationsschreiben.docx", "x")
    speichere(tmp_path / "verlauf" / "2026-09-18_motivationsschreiben.docx", "x")
    speichere(tmp_path / "Verlauf" / "2026-09-17_lebenslauf.docx", "x")
    speichere(tmp_path / "unter" / "tief" / "verlauf" / "alt.docx", "x")
    gefunden = finde_dateien(tmp_path)
    assert [p.relative_to(tmp_path).as_posix() for p in gefunden] == ["motivationsschreiben.docx"]


def test_durchsuche_zaehlt_verlauf_nicht_mit(tmp_path):
    speichere(tmp_path / "aktuell.docx", "[IBAN]")
    speichere(tmp_path / "verlauf" / "2026-09-18_aktuell.docx", "[IBAN]", "[AHV]")
    speichere(tmp_path / "unter" / "VERLAUF" / "ganz_alt.docx", "[IBAN]")
    e = durchsuche(tmp_path)
    assert e.dateien == ["aktuell.docx"]
    assert [(f.name, f.anzahl, f.dateien) for f in e.platzhalter] == [("IBAN", 1, ["aktuell.docx"])]


def test_durchsuche_lehnt_verlauf_als_gewaehlten_ordner_ab(tmp_path):
    speichere(tmp_path / "verlauf" / "2026-09-18_aktuell.docx", "[IBAN]")
    e = durchsuche(tmp_path / "verlauf")
    assert e.dateien == []
    assert e.platzhalter == []
    assert e.warnung is not None and "Verlauf" in e.warnung


def test_durchsuche_lehnt_auch_unterordner_im_verlauf_ab(tmp_path):
    speichere(tmp_path / "Verlauf" / "2026-09" / "alt.docx", "[IBAN]")
    e = durchsuche(tmp_path / "Verlauf" / "2026-09")
    assert e.dateien == []
    assert e.warnung is not None


def test_durchsuche_warnt_nicht_bei_normalem_ordner(tmp_path):
    speichere(tmp_path / "aktuell.docx", "[IBAN]")
    e = durchsuche(tmp_path)
    assert e.warnung is None
    assert e.dateien == ["aktuell.docx"]


def test_durchsuche_nimmt_pdf_formulare_mit(tmp_path):
    speichere(tmp_path / "brief.docx", "[ADRESSE]")
    formular(tmp_path / "antrag.pdf", {"Adresse": "[ADRESSE]", "Name": "[VORNAME NAME]", "Ort": "Bern"})
    e = durchsuche(tmp_path)
    assert e.dateien == ["antrag.pdf", "brief.docx"]
    assert [(f.name, f.anzahl, f.dateien) for f in e.platzhalter] == [
        ("ADRESSE", 2, ["antrag.pdf", "brief.docx"]),
        ("VORNAME NAME", 1, ["antrag.pdf"]),
    ]


def test_durchsuche_uebergeht_pdf_ohne_formular_und_mit_passwort(tmp_path):
    ohne_formular(tmp_path / "inserat.pdf")
    formular(tmp_path / "geschuetzt.pdf", {"A": "[AA]"}, passwort="geheim")
    e = durchsuche(tmp_path)
    assert e.dateien == []
    assert e.fehler == []


def test_durchsuche_meldet_kaputtes_pdf(tmp_path):
    (tmp_path / "kaputt.pdf").write_bytes(b"kein pdf")
    e = durchsuche(tmp_path)
    assert e.dateien == []
    assert len(e.fehler) == 1 and e.fehler[0].startswith("kaputt.pdf: ")
